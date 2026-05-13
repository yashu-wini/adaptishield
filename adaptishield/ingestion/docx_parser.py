# ingestion/docx_parser.py

import io
from pathlib import Path
from typing import Optional


class DocxParser:
    """Extracts text and metadata from .docx files using python-docx."""

    def parse_bytes(self, docx_bytes: bytes) -> dict:
        """Parse DOCX from raw bytes (for API uploads)."""
        from docx import Document
        doc = Document(io.BytesIO(docx_bytes))
        return self._extract(doc, source="upload")

    def parse(self, file_path: str) -> dict:
        from docx import Document
        path = Path(file_path)
        if not path.exists():
            raise FileNotFoundError(f"DOCX not found: {file_path}")

        doc = Document(file_path)
        return self._extract(doc, source=file_path)

    def _extract(self, doc, source: str = "unknown") -> dict:
        """Shared extraction logic for both file-path and bytes-based parsing."""
        paragraphs = []
        full_text = []

        for para in doc.paragraphs:
            if para.text.strip():
                paragraphs.append({
                    "style": para.style.name,
                    "text": para.text
                })
                full_text.append(para.text)

        # Extract tables
        tables_text = []
        for table in doc.tables:
            for row in table.rows:
                row_text = " | ".join(cell.text for cell in row.cells if cell.text.strip())
                if row_text:
                    tables_text.append(row_text)

        all_text = "\n".join(full_text + tables_text)

        props = doc.core_properties
        return {
            "text": all_text,
            "paragraphs": paragraphs,
            "metadata": {
                "title": props.title or "",
                "author": props.author or "",
                "source": source,
                "parser": "python-docx"
            }
        }
